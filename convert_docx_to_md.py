# -*- coding: utf-8 -*-
"""将技术方案 docx 转换为 Markdown"""
from docx import Document
import re

doc = Document(r'C:\Users\29144\Desktop\painting\FaceAngleLab_V3.1_单图AI头发重建技术方案.docx')

lines = []
in_code = False
code_buf = []

def flush_code():
    global in_code, code_buf
    if in_code:
        lines.append('```')
        lines.append('')
        in_code = False
        code_buf = []

for para in doc.paragraphs:
    text = para.text
    style = para.style.name if para.style else ''

    # 跳过空行
    if not text.strip():
        flush_code()
        lines.append('')
        continue

    # 标题
    if style.startswith('Heading 1'):
        flush_code()
        lines.append(f'# {text}')
        lines.append('')
    elif style.startswith('Heading 2'):
        flush_code()
        lines.append(f'## {text}')
        lines.append('')
    elif style.startswith('Heading 3'):
        flush_code()
        lines.append(f'### {text}')
        lines.append('')
    # 列表
    elif style.startswith('List Bullet'):
        flush_code()
        # 处理嵌套列表（List Bullet 2）
        indent = '  ' if '2' in style else ''
        lines.append(f'{indent}- {text}')
    elif style.startswith('List Number'):
        flush_code()
        lines.append(f'1. {text}')
    # 代码块（Consolas 字体）
    elif any(r.font.name == 'Consolas' for r in para.runs if r.font.name):
        if not in_code:
            lines.append('```python')
            in_code = True
        lines.append(text)
    # 普通段落
    else:
        flush_code()
        lines.append(text)
        lines.append('')

flush_code()

# 处理表格
md_lines = lines[:]
# 由于 docx 表格不在 paragraphs 中，需要单独处理
# 重新遍历 body 元素，按顺序插入表格
from docx.oxml.ns import qn

body = doc.element.body
result = []
para_idx = 0
table_idx = 0
tables = doc.tables

# 先收集所有段落文本（带样式）
all_paras = []
for para in doc.paragraphs:
    all_paras.append(para)

# 按文档顺序遍历
for child in body.iterchildren():
    if child.tag == qn('w:p'):
        if para_idx < len(md_lines):
            # 找到对应的 md 行
            # 简单处理：直接输出已有的 md_lines 中对应位置
            pass
        para_idx += 1
    elif child.tag == qn('w:tbl'):
        if table_idx < len(tables):
            table = tables[table_idx]
            result.append('')  # 表格前空行
            # 表头
            header = table.rows[0]
            headers = [cell.text.strip().replace('\n', ' ') for cell in header.cells]
            result.append('| ' + ' | '.join(headers) + ' |')
            result.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
            # 数据行
            for row in table.rows[1:]:
                cells = [cell.text.strip().replace('\n', '<br>') for cell in row.cells]
                result.append('| ' + ' | '.join(cells) + ' |')
            result.append('')
        table_idx += 1

# 更简单的方法：直接用之前生成的 lines，然后在适当位置插入表格
# 由于表格位置难以精确对应，我们采用另一种策略：
# 重新完整生成 markdown，按 body 顺序

print("正在按文档顺序重新生成...")

final = []
p_iter = iter(doc.paragraphs)
t_iter = iter(doc.tables)

para_list = list(doc.paragraphs)
table_list = list(doc.tables)
pi = 0
ti = 0

for child in body.iterchildren():
    if child.tag == qn('w:p'):
        if pi >= len(para_list):
            continue
        para = para_list[pi]
        pi += 1
        text = para.text
        style = para.style.name if para.style else ''

        if not text.strip():
            if in_code:
                # 代码块内的空行保留
                final.append('')
            else:
                final.append('')
            continue

        if style.startswith('Heading 1'):
            flush_code()
            final.append(f'# {text}')
            final.append('')
        elif style.startswith('Heading 2'):
            flush_code()
            final.append(f'## {text}')
            final.append('')
        elif style.startswith('Heading 3'):
            flush_code()
            final.append(f'### {text}')
            final.append('')
        elif style.startswith('List Bullet'):
            flush_code()
            indent = '  ' if '2' in style else ''
            final.append(f'{indent}- {text}')
        elif style.startswith('List Number'):
            flush_code()
            final.append(f'1. {text}')
        elif any(r.font.name == 'Consolas' for r in para.runs if r.font.name):
            if not in_code:
                final.append('```python')
                in_code = True
            final.append(text)
        else:
            flush_code()
            final.append(text)
            final.append('')

    elif child.tag == qn('w:tbl'):
        flush_code()
        if ti >= len(table_list):
            continue
        table = table_list[ti]
        ti += 1
        final.append('')
        header = table.rows[0]
        headers = [cell.text.strip().replace('\n', ' ') for cell in header.cells]
        final.append('| ' + ' | '.join(headers) + ' |')
        final.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
        for row in table.rows[1:]:
            cells = [cell.text.strip().replace('\n', '<br>') for cell in row.cells]
            final.append('| ' + ' | '.join(cells) + ' |')
        final.append('')

flush_code()

# 清理多余空行（连续3个以上空行压缩为2个）
cleaned = []
empty_count = 0
for line in final:
    if line.strip() == '':
        empty_count += 1
        if empty_count <= 2:
            cleaned.append(line)
    else:
        empty_count = 0
        cleaned.append(line)

md_content = '\n'.join(cleaned)

out_path = r'C:\Users\29144\Desktop\painting\FaceAngleLab_V3.1_单图AI头发重建技术方案.md'
with open(out_path, 'w', encoding='utf-8') as f:
    f.write(md_content)

print(f'Markdown 已生成：{out_path}')
print(f'总行数：{len(cleaned)}')
print(f'表格数：{len(table_list)}')
